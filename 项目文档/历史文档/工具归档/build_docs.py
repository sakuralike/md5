from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"E:\只比主题\minimax\密码侦探社\项目文档")
SRC = ROOT / "规范源文件"

NAVY = "111827"
PANEL = "1F2937"
AMBER = "F59E0B"
LIGHT = "F3F4F6"
GRAY = "6B7280"
PALE = "FFF7E6"
BORDER = "D1D5DB"
BLUE = "E8F3FF"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_run_font(run, name='Microsoft YaHei', size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name if name != 'Microsoft YaHei' else 'Aptos')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name if name != 'Microsoft YaHei' else 'Aptos')
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, field):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = field
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])


def add_rich_text(paragraph, text, default_size=10.5, default_color=None):
    # Parses **bold** and `inline code`.
    pattern = re.compile(r'(\*\*.*?\*\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r, size=default_size, color=default_color)
        token = m.group(0)
        if token.startswith('**'):
            r = paragraph.add_run(token[2:-2])
            set_run_font(r, size=default_size, bold=True, color=default_color)
        else:
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, name='Consolas', size=max(default_size-0.5, 8.5), color='7C2D12')
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, size=default_size, color=default_color)


def configure_document(doc, short_title):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ('Heading 1', 18, NAVY, 16, 8),
        ('Heading 2', 14, PANEL, 12, 6),
        ('Heading 3', 11.5, '92400E', 9, 4),
    ]:
        s = styles[name]
        s.font.name = 'Microsoft YaHei'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    styles['Heading 1'].paragraph_format.page_break_before = True

    for list_name in ['List Bullet', 'List Number']:
        s = styles[list_name]
        s.font.name = 'Microsoft YaHei'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(3)

    header = sec.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(short_title)
    set_run_font(r, size=8.5, color=GRAY)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), AMBER)
    pbdr.append(bottom)
    pPr.append(pbdr)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('密码侦探社 · ')
    set_run_font(r, size=8.5, color=GRAY)
    add_field(fp, 'PAGE')


def add_cover(doc, title, version, subtitle, status, date='2026-08-01'):
    # Hide header/footer on cover by using a distinct first page.
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('密码侦探社')
    set_run_font(r, size=16, bold=True, color=AMBER)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title)
    set_run_font(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(version)
    set_run_font(r, size=18, bold=True, color='92400E')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run(subtitle)
    set_run_font(r, size=11.5, color=GRAY)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.2)
    table.columns[1].width = Cm(8.5)
    for row, vals in zip(table.rows, [('文档状态', status), ('基线日期', date), ('维护方式', 'Markdown 源文件 + Word 评审版')]):
        for j, val in enumerate(vals):
            cell = row.cells[j]
            set_cell_margins(cell, 120, 150, 120, 150)
            set_cell_shading(cell, PALE if j == 0 else 'FFFFFF')
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=10.5, bold=(j == 0), color=NAVY if j == 0 else PANEL)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(45)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('授权恢复 · 本地验证 · 可审计共享')
    set_run_font(r, size=12, bold=True, color=AMBER)

    doc.add_page_break()


def extract_toc(lines):
    items=[]
    for line in lines:
        if line.startswith('## '):
            items.append((1, line[3:].strip()))
        elif line.startswith('### '):
            items.append((2, line[4:].strip()))
    return items


def add_toc(doc, lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('目录')
    set_run_font(r, size=20, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(12)
    for level, text in extract_toc(lines):
        if level == 2 and len(text) > 45:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7 if level == 1 else 1.5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        set_run_font(r, size=10.5 if level == 1 else 9.5, bold=(level == 1), color=PANEL if level == 1 else GRAY)
    doc.add_page_break()


def parse_table(lines, i):
    rows=[]
    while i < len(lines) and lines[i].strip().startswith('|'):
        parts=[p.strip() for p in lines[i].strip().strip('|').split('|')]
        if not all(re.fullmatch(r':?-{3,}:?', p or '') for p in parts):
            rows.append(parts)
        i += 1
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols=max(len(r) for r in rows)
    t=doc.add_table(rows=len(rows), cols=cols)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.style='Table Grid'
    t.autofit=True
    for ri,row in enumerate(rows):
        for ci in range(cols):
            val=row[ci] if ci < len(row) else ''
            cell=t.cell(ri,ci)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri==0:
                set_cell_shading(cell, NAVY)
            elif ri%2==0:
                set_cell_shading(cell, 'F9FAFB')
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            add_rich_text(p,val,default_size=8.8 if cols>=5 else 9.2,default_color='FFFFFF' if ri==0 else PANEL)
            if ri==0:
                for run in p.runs: run.bold=True
    set_repeat_table_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def add_code_block(doc, code_lines):
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Cm(0.4)
    p.paragraph_format.right_indent=Cm(0.4)
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(8)
    pPr=p._p.get_or_add_pPr()
    shd=OxmlElement('w:shd')
    shd.set(qn('w:fill'),'F3F4F6')
    pPr.append(shd)
    for idx,line in enumerate(code_lines):
        r=p.add_run(line)
        set_run_font(r,name='Consolas',size=8.2,color=PANEL)
        if idx < len(code_lines)-1:
            r.add_break()


def build(md_path, out_path, title, version, subtitle, status, short_title):
    raw=md_path.read_text(encoding='utf-8-sig')
    lines=raw.splitlines()
    doc=Document()
    configure_document(doc, short_title)
    add_cover(doc,title,version,subtitle,status)
    add_toc(doc,lines)

    i=0
    in_code=False
    code=[]
    first_h1_skipped=False
    while i<len(lines):
        line=lines[i].rstrip()
        stripped=line.strip()
        if stripped.startswith('```'):
            if not in_code:
                in_code=True; code=[]
            else:
                add_code_block(doc,code); in_code=False
            i+=1; continue
        if in_code:
            code.append(line); i+=1; continue
        if not stripped:
            i+=1; continue
        if stripped.startswith('|'):
            rows,i2=parse_table(lines,i)
            add_table(doc,rows); i=i2; continue
        if stripped.startswith('# '):
            if not first_h1_skipped:
                first_h1_skipped=True; i+=1; continue
            p=doc.add_paragraph(style='Heading 1'); add_rich_text(p,stripped[2:],default_size=18); i+=1; continue
        if stripped.startswith('## '):
            p=doc.add_paragraph(style='Heading 1'); add_rich_text(p,stripped[3:],default_size=18); i+=1; continue
        if stripped.startswith('### '):
            p=doc.add_paragraph(style='Heading 2'); add_rich_text(p,stripped[4:],default_size=14); i+=1; continue
        if stripped.startswith('#### '):
            p=doc.add_paragraph(style='Heading 3'); add_rich_text(p,stripped[5:],default_size=11.5); i+=1; continue
        if stripped.startswith('> '):
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Cm(0.7)
            p.paragraph_format.right_indent=Cm(0.5)
            pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),PALE); pPr.append(shd)
            add_rich_text(p,stripped[2:],default_size=10,color='92400E')
            i+=1; continue
        if re.match(r'^[-*] ', stripped):
            p=doc.add_paragraph(style='List Bullet')
            add_rich_text(p,stripped[2:])
            i+=1; continue
        if re.match(r'^\d+\. ', stripped):
            p=doc.add_paragraph(style='List Number')
            add_rich_text(p,re.sub(r'^\d+\.\s+','',stripped))
            i+=1; continue
        p=doc.add_paragraph()
        add_rich_text(p,stripped)
        i+=1

    # Add core properties
    doc.core_properties.title=f'{title} {version}'
    doc.core_properties.subject=subtitle
    doc.core_properties.author='密码侦探社项目组'
    doc.core_properties.keywords='密码侦探社, 项目规格, 开发计划, 安全, 压缩包恢复'
    doc.core_properties.comments='由原始项目文档整合、纠错和细化形成。'
    doc.save(out_path)
    print(out_path)

build(
    SRC/'密码侦探社项目规格说明书-v3.0.md',
    ROOT/'密码侦探社项目规格说明书-v3.0.docx',
    '项目规格说明书','v3.0',
    '统一产品定位、需求、业务规则、架构、安全与验收基线',
    '开发基线','项目规格说明书 v3.0')

build(
    SRC/'密码侦探社开发实施计划-v1.0.md',
    ROOT/'密码侦探社开发实施计划-v1.0.docx',
    '开发实施计划','v1.0',
    '14 周 MVP 路线图、任务拆解、里程碑、测试与发布门禁',
    '执行计划','开发实施计划 v1.0')
